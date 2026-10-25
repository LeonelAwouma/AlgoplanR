"""Délégué régional : Fait de Prince réservé au DRH, import d'une base nationale isolé sur la région, exports étanches.

Chaque test passe par l'API HTTP : les gardes de l'interface React n'interviennent pas.
"""
import io

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

from contextes.donnees.models import Enseignant, Etablissement, JeuDeDonnees
from domaine.acces.regions import REGIONS_CAMEROUN, region_canonique
from domaine.donnees.demonstration import construire_jeu_demonstration
from domaine.acces.isolement_regional import isoler_region
from tests.outils_acces import B, identifiant, session_complete

pytestmark = pytest.mark.django_db

J = "application/json"


def _national(nb_ecoles=100, nb_ens=600, graine=7):
    """Base nationale fictive : écoles réparties sur les dix régions, avec des graphies variées du nom."""
    e, n = construire_jeu_demonstration(graine, nb_ecoles, nb_ens)
    graphies = {"Extrême-Nord": "Extreme-Nord", "Nord-Ouest": "NORD OUEST", "Sud-Ouest": "sud-ouest"}
    for k, l in enumerate(e):
        region = REGIONS_CAMEROUN[k % len(REGIONS_CAMEROUN)]
        l["Region"] = graphies.get(region, region)
    return e, n


def _envoyer(client, etab, ens):
    ids = {}
    for typ, lignes in (("etablissements", etab), ("enseignants", ens)):
        buf = io.BytesIO()
        pd.DataFrame(lignes).to_excel(buf, index=False)
        buf.name = f"{typ}_national.xlsx"
        buf.seek(0)
        r = client.post(B + "fichiers/", {"type": typ, "fichier": buf})
        assert r.status_code == 201, r.content
        ids[typ] = r.json()["id"]
    return ids


def _importer(client, etab, ens, nom="Import"):
    ids = _envoyer(client, etab, ens)
    return client.post(B + "jeux/", {"nom": nom, "fichier_etablissements": ids["etablissements"],
                                     "fichier_enseignants": ids["enseignants"]}, content_type=J)


def _attendu(etab, ens, region):
    ecoles = {l["Identifiant_Ecole"] for l in etab if region_canonique(l["Region"]) == region}
    return ecoles, {l["Matricule_Enseignant"] for l in ens if l["Ecole_Attache_Actuelle"] in ecoles}


# --- Fait de Prince : DRH seulement -------------------------------------------------------------------------------

def _routes_prince(jeu):
    base = f"jeux/{jeu}/faits-prince/"
    return [("get", base), ("post", base), ("post", base + "impact/"), ("post", base + "import/"),
            ("get", base + "application/"), ("get", base + "bases/etablissements/"),
            ("get", base + "bases/enseignants/export/"), ("delete", base + "1/")]


def test_delegue_ne_peut_utiliser_aucune_route_du_fait_de_prince(centre, littoral, drh):
    jeu = centre.post(B + "jeux/demonstration/", {"nb_ecoles": 12, "nb_enseignants": 60},
                      content_type=J).json()["id"]
    for client in (centre, littoral):   # son propre jeu comme celui d'un autre : refus avant toute lecture du jeu
        for methode, url in _routes_prince(jeu):
            r = getattr(client, methode)(B + url, **({"data": {}, "content_type": J} if methode == "post" else {}))
            assert r.status_code == 403 and r.json()["code"] == "reserve_drh", (methode, url, r.status_code)
    # Même un code d'accès Effet de Prince ne change rien : le contrôle est celui du compte.
    r = centre.post(B + f"jeux/{jeu}/faits-prince/", {"id_ens": "M1", "id_etab": "E1"}, content_type=J,
                    HTTP_X_CODE_PRINCE="n'importe-quoi")
    assert r.status_code == 403 and r.json()["code"] == "reserve_drh"
    assert JeuDeDonnees.objects.get(pk=jeu).faits_prince.count() == 0


def test_delegue_ne_joint_pas_de_fichier_du_prince_a_la_creation(centre):
    e, n = _national(20, 80)
    for l in e:
        l["Region"] = "Centre"
    ids = _envoyer(centre, e, n)
    buf = io.BytesIO(b"id_ens,id_etab\nM1,ECO001\n")
    buf.name = "prince.csv"
    prince = centre.post(B + "fichiers/", {"type": "prince", "fichier": buf})
    r = centre.post(B + "jeux/", {"fichier_etablissements": ids["etablissements"],
                                  "fichier_enseignants": ids["enseignants"],
                                  "fichier_prince": prince.json()["id"]}, content_type=J)
    assert r.status_code in (403, 422)
    assert JeuDeDonnees.objects.count() == 0


def test_drh_conserve_le_fait_de_prince(drh):
    jeu = drh.post(B + "jeux/demonstration/", {"nb_ecoles": 12, "nb_enseignants": 60}, content_type=J).json()["id"]
    assert drh.get(B + f"jeux/{jeu}/faits-prince/").status_code == 200
    assert drh.get(B + f"jeux/{jeu}/faits-prince/application/").status_code == 200
    assert drh.get(B + f"jeux/{jeu}/faits-prince/bases/etablissements/").status_code == 200
    assert drh.post(B + f"jeux/{jeu}/faits-prince/impact/", {"id_ens": "?", "id_etab": "?"},
                    content_type=J).status_code != 403


# --- Modes Décideur et Analyste : aucune permission ne dépend de l'affichage --------------------------------------

def test_les_routes_du_mode_analyste_servent_le_delegue_sans_elargir_son_perimetre(centre, littoral, drh):
    jeu = centre.post(B + "jeux/demonstration/", {"nb_ecoles": 20, "nb_enseignants": 120},
                      content_type=J).json()["id"]
    ecole = centre.get(B + f"jeux/{jeu}/etablissements/").json()["resultats"][0]["Identifiant_Ecole"]
    ens = centre.get(B + f"jeux/{jeu}/enseignants/").json()["resultats"][0]["Matricule_Enseignant"]
    analyste = [f"jeux/{jeu}/etablissements/{ecole}/", f"jeux/{jeu}/enseignants/{ens}/", f"jeux/{jeu}/diagnostic/",
                f"jeux/{jeu}/qualite/", f"jeux/{jeu}/scenarios/", f"jeux/{jeu}/rapports/territorial/",
                f"jeux/{jeu}/hypotheses/"]
    for url in analyste:
        assert centre.get(B + url).status_code == 200, url
        assert littoral.get(B + url).status_code == 404, url      # accès inter-régional refusé
    # Le DRH lit toutes les régions, y compris pour ces routes détaillées.
    assert all(drh.get(B + url).status_code == 200 for url in analyste)


# --- Import d'une base nationale par un délégué --------------------------------------------------------------------

@pytest.mark.parametrize("region", ["Centre", "Littoral", "Extrême-Nord"])
def test_delegue_importe_la_base_nationale_et_ne_garde_que_sa_region(poste_active, region):
    etab, ens = _national()
    assert {region_canonique(l["Region"]) for l in etab} == set(REGIONS_CAMEROUN)   # les dix régions
    client = session_complete(identifiant(region), region)
    r = _importer(client, etab, ens)
    assert r.status_code == 201, r.content
    jeu = r.json()
    ecoles, matricules = _attendu(etab, ens, region)
    assert jeu["portee"] == "regional" and jeu["region"] == region
    assert (jeu["nb_etablissements"], jeu["nb_enseignants"]) == (len(ecoles), len(matricules))
    # Rien d'une autre région n'est stocké.
    en_base = JeuDeDonnees.objects.get(pk=jeu["id"])
    assert {e.id_ecole for e in Etablissement.objects.filter(jeu=en_base)} == ecoles
    assert {t.matricule for t in Enseignant.objects.filter(jeu=en_base)} == matricules
    assert {e.region for e in Etablissement.objects.filter(jeu=en_base)} == {region}
    # Bilan de l'import, dans le rapport de qualité.
    bilan = jeu["qualite"]["import_regional"]
    assert bilan["etablissements_fichier"] == len(etab) and bilan["etablissements_retenus"] == len(ecoles)
    assert bilan["enseignants_fichier"] == len(ens) and bilan["enseignants_retenus"] == len(matricules)
    assert bilan["etablissements_autres_regions"] == len(etab) - len(ecoles)
    assert bilan["enseignants_autres_regions"] == len(ens) - len(matricules)
    assert bilan["enseignants_non_rattaches"] == bilan["enseignants_ambigus"] == 0
    # Effectifs et statistiques : ceux de la région seulement.
    assert client.get(B + f"jeux/{jeu['id']}/etablissements/").json()["count"] == len(ecoles)
    assert client.get(B + f"jeux/{jeu['id']}/enseignants/").json()["count"] == len(matricules)


def test_la_region_demandee_par_l_interface_est_ignoree(centre, littoral):
    etab, ens = _national()
    jeu = _importer(centre, etab, ens).json()["id"]
    ecoles, _ = _attendu(etab, ens, "Centre")
    for param in ("region=Littoral", "region=Littoral&departement=Wouri", "region=%2A"):
        r = centre.get(B + f"jeux/{jeu}/etablissements/?{param}&taille=500")
        assert r.status_code in (200, 403), param
        if r.status_code == 200:
            assert {l["Identifiant_Ecole"] for l in r.json()["resultats"]} <= ecoles
    # Un champ « region » ajouté à la création n'est pas lu : la région vient de la session.
    ids = _envoyer(littoral, etab, ens)
    r = littoral.post(B + "jeux/", {"fichier_etablissements": ids["etablissements"],
                                    "fichier_enseignants": ids["enseignants"], "region": "Centre", "portee": "national"},
                      content_type=J)
    assert r.status_code == 201 and r.json()["region"] == "Littoral" and r.json()["portee"] == "regional"


def test_l_import_d_un_delegue_n_altere_aucun_autre_jeu(drh, centre, littoral):
    etab, ens = _national()
    national = _importer(drh, etab, ens, "National").json()
    assert national["portee"] == "national" and national["nb_etablissements"] == len(etab)
    jeu_littoral = _importer(littoral, etab, ens, "Littoral").json()
    avant = {j.pk: (j.etablissements.count(), j.enseignants.count(), j.nom) for j in JeuDeDonnees.objects.all()}
    jeu_centre = _importer(centre, etab, ens, "Centre").json()
    for pk, empreinte in avant.items():
        j = JeuDeDonnees.objects.get(pk=pk)
        assert (j.etablissements.count(), j.enseignants.count(), j.nom) == empreinte
    assert JeuDeDonnees.objects.count() == len(avant) + 1
    # Le délégué ne voit ni le jeu national ni celui d'une autre région ; le DRH voit les trois.
    assert centre.get(B + f"jeux/{national['id']}/").status_code == 404
    assert centre.get(B + f"jeux/{jeu_littoral['id']}/").status_code == 404
    assert {j["id"] for j in centre.get(B + "jeux/").json()} == {jeu_centre["id"]}
    assert {j["id"] for j in drh.get(B + "jeux/").json()} == {national["id"], jeu_littoral["id"], jeu_centre["id"]}
    # Le DRH ne modifie pas un jeu régional.
    assert drh.delete(B + f"jeux/{jeu_centre['id']}/").status_code == 403


def test_le_drh_garde_l_import_national_et_les_dix_regions(drh):
    etab, ens = _national()
    jeu = _importer(drh, etab, ens).json()
    assert (jeu["nb_etablissements"], jeu["nb_enseignants"]) == (len(etab), len(ens))
    assert "import_regional" not in jeu["qualite"]
    regions = {l["Region"] for l in drh.get(B + f"jeux/{jeu['id']}/etablissements/?taille=500").json()["resultats"]}
    assert regions == set(REGIONS_CAMEROUN)      # graphies ramenées au nom officiel
    for region in REGIONS_CAMEROUN:
        r = drh.get(B + f"jeux/{jeu['id']}/etablissements/", {"region": region})
        assert r.status_code == 200 and r.json()["count"] == sum(1 for l in etab if region_canonique(l["Region"]) == region)


def test_le_fichier_source_n_est_pas_modifie(centre):
    etab, ens = _national(40, 200)
    ids = _envoyer(centre, etab, ens)
    from contextes.donnees.models import FichierImporte
    avant = {f.pk: (len(f.lignes), f.lignes[0]) for f in FichierImporte.objects.all()}
    assert centre.post(B + "jeux/", {"fichier_etablissements": ids["etablissements"],
                                     "fichier_enseignants": ids["enseignants"]}, content_type=J).status_code == 201
    assert {f.pk: (len(f.lignes), f.lignes[0]) for f in FichierImporte.objects.all()} == avant


# --- Rattachements douteux : signalés, jamais attribués arbitrairement ------------------------------------------------

def test_enseignants_non_rattaches_ou_ambigus_sont_signales_et_exclus(centre):
    etab, ens = _national()
    ecoles_centre = sorted(_attendu(etab, ens, "Centre")[0])
    autre = next(l for l in etab if region_canonique(l["Region"]) == "Littoral")
    # Libellés de région : casse, accents, espaces.
    for l in etab:
        if l["Identifiant_Ecole"] == ecoles_centre[0]:
            l["Region"] = "  CENTRE "
    # 1. enseignant sans école ; 2. école absente de la base ; 3. école dont le code est porté par deux régions.
    ens[0] = {**ens[0], "Ecole_Attache_Actuelle": ""}
    ens[1] = {**ens[1], "Ecole_Attache_Actuelle": "INCONNUE-999"}
    doublon = {**autre, "Region": "Centre"}                   # même code, région différente de la ligne d'origine
    etab.append(doublon)
    ens[2] = {**ens[2], "Ecole_Attache_Actuelle": autre["Identifiant_Ecole"]}
    # 4. région de l'enseignant contradictoire avec celle de son école.
    for l in ens:
        l["Region"] = ""
    ens[3] = {**ens[3], "Ecole_Attache_Actuelle": ecoles_centre[1], "Region": "Littoral"}
    # 5. école du Centre sans région lisible.
    etab.append({**etab[0], "Identifiant_Ecole": "ECO-SANS-REGION", "Region": "Atlantis"})
    r = _importer(centre, etab, ens)
    assert r.status_code == 201, r.content
    jeu = r.json()
    bilan = jeu["qualite"]["import_regional"]
    assert bilan["enseignants_non_rattaches"] == 2
    # Deux enseignants contradictoires/ambigus de la liste, plus tous ceux du code d'école ambigu.
    assert bilan["enseignants_ambigus"] >= 2
    ambigus_attaches = sum(1 for l in ens if l["Ecole_Attache_Actuelle"] == autre["Identifiant_Ecole"])
    assert bilan["enseignants_ambigus"] == ambigus_attaches + 1
    assert bilan["etablissements_ecartes"] >= 3                # deux lignes du code ambigu + région illisible
    codes = {a["code"] for a in bilan["anomalies"]}
    assert {"sans_ecole", "ecole_inconnue", "ecole_indeterminee", "region_contradictoire", "code_ambigu",
            "region_illisible"} <= codes
    en_base = JeuDeDonnees.objects.get(pk=jeu["id"])
    matricules = {t.matricule for t in en_base.enseignants.all()}
    assert not matricules & {ens[i]["Matricule_Enseignant"] for i in range(4)}
    assert autre["Identifiant_Ecole"] not in {e.id_ecole for e in en_base.etablissements.all()}
    assert "ECO-SANS-REGION" not in {e.id_ecole for e in en_base.etablissements.all()}
    assert ecoles_centre[0] in {e.id_ecole for e in en_base.etablissements.all()}      # « CENTRE » normalisé
    avert = {a["code"] for a in jeu["qualite"]["avertissements"]}
    assert {"ENS_NON_RATTACHES", "ENS_AMBIGUS", "ETAB_ECARTES"} <= avert


def test_isolement_par_code_jamais_par_nom():
    etab = [{"Identifiant_Ecole": "A1", "Region": "Centre", "nom_etab": "École Nord"},
            {"Identifiant_Ecole": "B1", "Region": "Nord", "nom_etab": "École Nord"}]
    ens = [{"Ecole_Attache_Actuelle": "B1"}, {"Ecole_Attache_Actuelle": "École Nord"}]
    e, n, bilan = isoler_region("Centre", etab, ens, None, None, col_region="Region", col_id="Identifiant_Ecole",
                                col_ecole="Ecole_Attache_Actuelle")
    assert (e, n) == ([0], [])
    assert bilan["enseignants_autres_regions"] == 1 and bilan["enseignants_non_rattaches"] == 1


# --- Exports d'un délégué ----------------------------------------------------------------------------------------------

def _textes_xlsx(octets):
    wb = load_workbook(io.BytesIO(octets))
    return [str(c.value).strip() for ws in wb for ligne in ws.iter_rows() for c in ligne if c.value is not None]


def _textes_docx(octets):
    d = Document(io.BytesIO(octets))
    textes = [p.text.strip() for p in d.paragraphs]
    for t in d.tables:
        for ligne in t.rows:
            textes += [c.text.strip() for c in ligne.cells]
    return textes


def test_exports_excel_et_word_du_delegue_sans_donnee_d_une_autre_region(centre):
    etab, ens = _national()
    jeu = _importer(centre, etab, ens).json()["id"]
    ecoles_centre, _ = _attendu(etab, ens, "Centre")
    autres_ecoles = {l["Identifiant_Ecole"] for l in etab} - ecoles_centre
    autres_regions = set(REGIONS_CAMEROUN) - {"Centre"}
    for fmt, lire in (("xlsx", _textes_xlsx), ("docx", _textes_docx)):
        for params in ({}, {"region": "Littoral"}, {"niveau": "ecole"}):
            r = centre.get(B + f"jeux/{jeu}/rapports/territorial/export/", {"format": fmt, **params})
            if r.status_code != 200:
                assert r.status_code in (403, 404, 422), (fmt, params, r.status_code)
                continue
            textes = lire(r.content)
            assert not set(textes) & autres_regions, (fmt, params)
            assert not any(i in t for t in textes for i in autres_ecoles), (fmt, params)
    # L'export « national complet » est réservé au DRH.
    r = centre.get(B + f"jeux/{jeu}/rapports/territorial/export/", {"format": "xlsx", "type": "national_complet"})
    assert r.status_code == 403
    # Les exports du rapport JSON ne contiennent que la région.
    rapport = centre.get(B + f"jeux/{jeu}/rapports/territorial/", {"region": "Littoral"})
    assert rapport.status_code in (200, 403)
    if rapport.status_code == 200:
        assert not any(r in rapport.content.decode() for r in ("Littoral", "Extrême-Nord"))


# --- Scénario retenu : explicite, borné au jeu du délégué ------------------------------------------------------------

def test_scenario_retenu_explicite_et_base_de_calcul_indiquee(drh, centre, littoral):
    jeu = centre.post(B + "jeux/demonstration/", {"nb_ecoles": 20, "nb_enseignants": 120}, content_type=J).json()
    autre = littoral.post(B + "jeux/demonstration/", {"nb_ecoles": 20, "nb_enseignants": 120}, content_type=J).json()
    # Aucun scénario n'est retenu automatiquement : le serveur dit sur quoi il calcule, et que c'est provisoire.
    assert jeu["scenario_retenu_id"] is None
    base = centre.get(B + f"jeux/{jeu['id']}/planification/etat/").json()["base_calcul"]
    assert base["type"] == "scenario_par_defaut" and base["provisoire"] is True and base["scenario_nom"]
    scenarios = centre.get(B + f"jeux/{jeu['id']}/scenarios/").json()
    choisi = next(s for s in scenarios if s["calcule_le"])["id"]
    assert choisi
    # Un délégué retient un scénario de SON jeu ; le jeu d'une autre région et le jeu du DRH ne bougent pas.
    r = centre.put(B + f"jeux/{jeu['id']}/scenario-retenu/", {"scenario": choisi}, content_type=J)
    assert r.status_code == 200 and r.json()["scenario_retenu_id"] == str(choisi)
    base = centre.get(B + f"jeux/{jeu['id']}/planification/etat/").json()["base_calcul"]
    assert base["type"] == "scenario_retenu" and base["provisoire"] is False
    assert littoral.get(B + f"jeux/{autre['id']}/").json()["scenario_retenu_id"] is None
    assert littoral.put(B + f"jeux/{jeu['id']}/scenario-retenu/", {"scenario": choisi},
                        content_type=J).status_code == 404
    assert drh.put(B + f"jeux/{jeu['id']}/scenario-retenu/", {"scenario": choisi}, content_type=J).status_code == 403
