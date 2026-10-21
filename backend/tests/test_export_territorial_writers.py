"""Les écrivains xlsx et docx du rapport territorial (domaine pur, sans Django/BD) : mise en forme, en-têtes
répétés, neutralisation des formules injectées, méthodologie."""
import io

from openpyxl import load_workbook
from docx import Document

from contextes.rapports import export_docx as W
from contextes.rapports import export_xlsx as X

REGISTRE = [{"code": "encadrement", "intitule": "Taux d'encadrement", "formule": "f", "population": "p",
            "source": "s", "limites": "l"},
           {"code": "alea", "intitule": "Degré d'aléa", "formule": "f2", "population": "p2", "source": "s2",
            "limites": "l2"}]


def _rapport(titre="Cameroun", scenario_id="7"):
    return {"titre": titre, "scenario_id": scenario_id, "registre_indicateurs": REGISTRE}


def _tableaux(territoire="Centre"):
    return [
        {"code": "encadrement", "intitule": "Taux d'encadrement", "formule": "f", "population": "p", "source": "s",
         "limites": "l", "colonnes": ["Import", "Après simulation"],
         "lignes": [{"territoire": "Cameroun", "valeurs": [52.3, 48.1]},
                    {"territoire": territoire, "valeurs": [40.0, "Sans objet"]}]},
        {"code": "alea", "intitule": "Degré d'aléa", "formule": "f2", "population": "p2", "source": "s2",
         "limites": "l2", "colonnes": ["Import"],
         "lignes": [{"territoire": "Cameroun", "valeurs": ["Non calculable"]}]},
    ]


def test_classeur_perimetre_feuilles_et_methodologie():
    octets = X.classeur_perimetre(_rapport(), _tableaux())
    assert octets[:2] == b"PK"
    wb = load_workbook(io.BytesIO(octets))
    assert set(wb.sheetnames) == {"Taux d'encadrement", "Degré d'aléa", "Méthodologie"}
    ws = wb["Taux d'encadrement"]
    assert [c.value for c in ws[1]] == ["Territoire", "Import", "Après simulation"]
    assert [c.value for c in ws[2]] == ["Cameroun", 52.3, 48.1]
    assert [c.value for c in ws[3]] == ["Centre", 40.0, "Sans objet"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:C1"
    meth = wb["Méthodologie"]
    assert [c.value for c in meth[1]] == ["Code", "Intitulé", "Formule", "Population", "Source", "Limites"]
    assert meth.cell(row=2, column=1).value == "encadrement"


def test_classeur_neutralise_une_formule_injectee_dans_un_nom_de_territoire():
    octets = X.classeur_perimetre(_rapport(), _tableaux(territoire="=CMD|' /C calc'!A1"))
    wb = load_workbook(io.BytesIO(octets))
    ws = wb["Taux d'encadrement"]
    valeur = ws.cell(row=3, column=1).value
    assert not str(valeur).startswith("=")   # neutralisée (préfixée), jamais exécutable à l'ouverture


def test_document_perimetre_titre_tableaux_et_notes():
    octets = W.document_perimetre(_rapport(), _tableaux(), jeu_nom="Jeu de test")
    assert octets[:2] == b"PK"
    doc = Document(io.BytesIO(octets))
    textes = [p.text for p in doc.paragraphs]
    assert "Cameroun" in textes[0]
    assert any("Jeu de test" in t and "7" in t for t in textes)
    assert any(t == "Taux d'encadrement" for t in textes)
    assert any("Formule : f." in t and "Limites : l." in t for t in textes)
    assert len(doc.tables) == 2
    t0 = doc.tables[0]
    assert [c.text for c in t0.rows[0].cells] == ["Territoire", "Import", "Après simulation"]
    assert [c.text for c in t0.rows[2].cells] == ["Centre", "40", "Sans objet"]
    # Page paysage (les tableaux comptent jusqu'à 8 colonnes de situations).
    section = doc.sections[0]
    assert section.page_width > section.page_height


def test_document_entete_de_tableau_repete():
    octets = W.document_perimetre(_rapport(), _tableaux())
    doc = Document(io.BytesIO(octets))
    from docx.oxml.ns import qn
    tr_pr = doc.tables[0].rows[0]._tr.find(qn("w:trPr"))
    assert tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None


def test_document_neutralise_une_formule_injectee():
    octets = W.document_perimetre(_rapport(), _tableaux(territoire="=1+1"))
    doc = Document(io.BytesIO(octets))
    cellule = doc.tables[0].rows[2].cells[0].text
    assert not cellule.startswith("=")


def test_document_national_complet_sections_et_sauts_de_page():
    complet = {
        "national": {"titre": "Cameroun", "niveau": "national", "tableaux": _tableaux()},
        "regions": [
            {"titre": "Centre", "niveau": "region", "tableaux": _tableaux("Mfoundi"),
             "departements": [{"titre": "Centre > Mfoundi", "niveau": "departement",
                               "tableaux": _tableaux("Yaoundé 1")}]},
            {"titre": "Littoral", "niveau": "region", "tableaux": _tableaux("Wouri"), "departements": []},
        ],
        "registre_indicateurs": REGISTRE,
    }
    octets = W.document_national_complet(complet, jeu_nom="Jeu", scenario_id="1")
    doc = Document(io.BytesIO(octets))
    titre_document = next(p.text for p in doc.paragraphs if p.style.name == "Title")
    assert titre_document == "Rapport national complet"
    titres = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Centre" in titres and "Littoral" in titres and "Centre > Mfoundi" in titres
    # national (2 tableaux) + région Centre (2) + département Mfoundi (2) + région Littoral (2) = 8 tableaux.
    assert len(doc.tables) == 8


def test_classeur_national_complet_une_feuille_nationale_une_par_region():
    complet = {
        "national": {"titre": "Cameroun", "niveau": "national", "tableaux": _tableaux()},
        "regions": [{"titre": "Centre", "niveau": "region", "tableaux": _tableaux("Mfoundi"), "departements": []}],
        "registre_indicateurs": REGISTRE,
    }
    octets = X.classeur_national_complet(complet)
    wb = load_workbook(io.BytesIO(octets))
    assert "Méthodologie" in wb.sheetnames
    assert len(wb.sheetnames) == 2 * 2 + 1   # (national + 1 région) × 2 indicateurs + méthodologie
