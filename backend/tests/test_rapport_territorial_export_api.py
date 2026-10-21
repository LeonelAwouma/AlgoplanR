"""API d'export du rapport territorial (sous-lot 4.7) : GET jeux/{jeu}/rapports/territorial/export/."""
import io

import pytest
from docx import Document
from openpyxl import load_workbook

pytestmark = pytest.mark.django_db
B = "/api/v1/"


@pytest.fixture
def c(drh):
    return drh


@pytest.fixture
def jeu(c):
    r = c.post(B + "jeux/demonstration/", {"nb_ecoles": 30, "nb_enseignants": 220}, content_type="application/json")
    assert r.status_code == 201, r.content
    return r.json()["id"]


def _get_fichier(c, url, **kw):
    r = c.get(B + url, kw)
    assert r.status_code == 200, (url, r.status_code, r.content[:500])
    assert r["Content-Disposition"].startswith("attachment")
    return r.content


def test_xlsx_national_coherent_avec_le_json(c, jeu):
    r = c.get(B + f"jeux/{jeu}/rapports/territorial/")
    rapport = r.json()
    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="xlsx")
    assert octets[:2] == b"PK"
    wb = load_workbook(io.BytesIO(octets))
    assert "Méthodologie" in wb.sheetnames
    ws = wb["Taux d'encadrement"]
    entetes = [cell.value for cell in ws[1]]
    assert entetes[0] == "Territoire"
    ligne_cameroun = [cell.value for cell in ws[2]]
    assert ligne_cameroun[0] == rapport["titre"]
    taux_json = rapport["noeud"]["encadrement"]["apres_simulation"]["taux"]
    # Unité « élèves par enseignant » dans l'en-tête : jamais un pourcentage.
    idx = entetes.index("Après simulation (élèves par enseignant)")
    assert ligne_cameroun[idx] == taux_json
    # Exhaustivité : une ligne par région, sans pagination.
    assert ws.max_row - 1 == 1 + len(rapport["enfants"])


def test_docx_national_ouvre_et_correspond_au_json(c, jeu):
    r = c.get(B + f"jeux/{jeu}/rapports/territorial/")
    rapport = r.json()
    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="docx")
    assert octets[:2] == b"PK"
    doc = Document(io.BytesIO(octets))   # ouverture réelle, pas seulement un contrôle de zip
    assert doc.tables
    table = doc.tables[0]
    assert [c0.text for c0 in table.rows[0].cells][0] == "Territoire"
    assert [c0.text for c0 in table.rows[1].cells][0] == rapport["titre"]


def test_filtre_indicateur_limite_a_une_feuille_et_une_table(c, jeu):
    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="xlsx", indicateur="alea")
    wb = load_workbook(io.BytesIO(octets))
    assert set(wb.sheetnames) == {"Degré d'aléa", "Méthodologie"}

    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="docx", indicateur="alea")
    doc = Document(io.BytesIO(octets))
    assert len(doc.tables) == 1


def test_export_respecte_la_selection_regionale(c, jeu):
    national = c.get(B + f"jeux/{jeu}/rapports/territorial/").json()
    region = national["enfants"][0]["territoire"]["region"]
    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="xlsx", region=region)
    wb = load_workbook(io.BytesIO(octets))
    ws = wb["Taux d'encadrement"]
    assert ws.cell(row=2, column=1).value == region


def test_national_complet_drh_seulement(c, jeu, centre):
    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="docx", type="national_complet")
    doc = Document(io.BytesIO(octets))
    titre = next(p.text for p in doc.paragraphs if p.style.name == "Title")
    assert titre == "Rapport national complet"
    assert len(doc.sections) >= 1

    jeu_centre = centre.post(B + "jeux/demonstration/", {"nb_ecoles": 10, "nb_enseignants": 70},
                             content_type="application/json").json()["id"]
    r = centre.get(B + f"jeux/{jeu_centre}/rapports/territorial/export/", {"format": "xlsx", "type": "national_complet"})
    assert r.status_code == 403, r.content


def test_national_complet_xlsx_exhaustif(c, jeu):
    """Décision corrigée : le classeur national complet est exhaustif (régions, départements, arrondissements,
    établissements), pas seulement synthétique — vérifié par le nombre de lignes de chaque feuille."""
    from contextes.donnees import depots
    from contextes.donnees.models import JeuDeDonnees

    ecoles = depots.charger_ecoles(JeuDeDonnees.objects.get(id=jeu))
    regions = {e.region for e in ecoles if e.region}
    departements = {(e.region, e.departement) for e in ecoles if e.region and e.departement}
    communes = {(e.region, e.departement, e.commune) for e in ecoles if e.region and e.departement and e.commune}

    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="xlsx", type="national_complet")
    wb = load_workbook(io.BytesIO(octets))
    for nom in ("Régions", "Départements", "Arrondissements", "Établissements"):
        assert nom in wb.sheetnames, wb.sheetnames

    assert wb["Régions"].max_row - 1 == len(regions)
    assert wb["Départements"].max_row - 1 == len(departements)
    assert wb["Arrondissements"].max_row - 1 == len(communes)
    assert wb["Établissements"].max_row - 1 == len(ecoles)
    assert wb["Établissements"].auto_filter.ref is not None

    # Cohérence avec le JSON national (4.6) : même taux d'encadrement après simulation pour une région donnée.
    national = c.get(B + f"jeux/{jeu}/rapports/territorial/").json()
    noms_json = {e["territoire"]["nom"]: e["encadrement"]["apres_simulation"]["taux"] for e in national["enfants"]}
    ws = wb["Régions"]
    entetes = [cell.value for cell in ws[1]]
    idx_territoire, idx_encadrement = entetes.index("Territoire"), entetes.index("Taux d'encadrement (élèves par enseignant)")
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert row[idx_encadrement] == noms_json[row[idx_territoire]]


def test_format_et_type_invalides_refuses(c, jeu):
    r = c.get(B + f"jeux/{jeu}/rapports/territorial/export/", {"format": "pdf"})
    assert r.status_code == 422, r.content
    r = c.get(B + f"jeux/{jeu}/rapports/territorial/export/", {"format": "xlsx", "type": "autre"})
    assert r.status_code == 422, r.content
    r = c.get(B + f"jeux/{jeu}/rapports/territorial/export/", {"format": "xlsx", "indicateur": "inconnu"})
    assert r.status_code == 422, r.content


def test_delegue_restreint_a_sa_region(centre):
    jeu = centre.post(B + "jeux/demonstration/", {"nb_ecoles": 15, "nb_enseignants": 100},
                      content_type="application/json").json()["id"]
    octets = _get_fichier(centre, f"jeux/{jeu}/rapports/territorial/export/", format="xlsx", region="Littoral")
    wb = load_workbook(io.BytesIO(octets))
    ws = wb["Taux d'encadrement"]
    assert ws.cell(row=2, column=1).value == "Centre"


def test_sans_session_refuse(jeu):
    from django.test import Client

    r = Client().get(B + f"jeux/{jeu}/rapports/territorial/export/", {"format": "xlsx"})
    assert r.status_code == 401


def test_export_niveau_et_annee_choisis_sont_respectes(c, jeu):
    """Le panneau d'export envoie ventilation (`niveau`), indicateur et année : le fichier suit la sélection."""
    json_communes = c.get(B + f"jeux/{jeu}/rapports/territorial/", {"niveau": "commune"}).json()
    octets = _get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", format="xlsx", niveau="commune",
                          indicateur="encadrement", annee="n2")
    wb = load_workbook(io.BytesIO(octets))
    ws = wb["Taux d'encadrement"]
    entetes = [cell.value for cell in ws[1]]
    assert "N+2 (élèves par enseignant)" in entetes
    assert not any(str(e).startswith(("N+1", "N+3", "N+4", "N+5")) for e in entetes)
    assert ws.max_row - 1 == 1 + len(json_communes["enfants"])
    assert wb.sheetnames == ["Taux d'encadrement", "Méthodologie"]


@pytest.mark.parametrize("param,valeur", [("annee", "n9"), ("niveau", "planete")])
def test_export_refuse_un_niveau_ou_une_annee_inconnus(c, jeu, param, valeur):
    r = c.get(B + f"jeux/{jeu}/rapports/territorial/export/", {"format": "xlsx", param: valeur})
    assert r.status_code == 422


def test_l_unite_du_taux_d_encadrement_n_est_jamais_un_pourcentage(c, jeu):
    registre = {d["code"]: d for d in c.get(B + f"jeux/{jeu}/rapports/territorial/").json()["registre_indicateurs"]}
    assert registre["encadrement"]["unite"] == "élèves par enseignant"
    assert all("%" in registre[k]["unite"] for k in ("couverture_globale", "couverture_operationnelle", "couverture_parec"))


def test_word_est_en_a4_paysage_avec_scenario_et_decimale_francaise(c, jeu):
    for params in ({"format": "docx"}, {"format": "docx", "type": "national_complet"}):
        doc = Document(io.BytesIO(_get_fichier(c, f"jeux/{jeu}/rapports/territorial/export/", **params)))
        section = doc.sections[0]
        assert round(section.page_width.cm, 1) == 29.7 and round(section.page_height.cm, 1) == 21.0   # A4 paysage
        sous_titre = next(p.text for p in doc.paragraphs if p.text.startswith("Jeu :"))
        assert "Scénario : " in sous_titre and "Scénario :  " not in sous_titre and "Scénario : ·" not in sous_titre
        cellules = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        assert not any(x.replace(".", "", 1).isdigit() and "." in x for x in cellules), "décimale avec point dans le Word"
